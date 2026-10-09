"""Executive Brief DOCX Generator matching ThaiNDC executive format."""
import io, re
from typing import Optional
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

FONT_NAME = 'TH SarabunPSK'

def set_run_font(run, size_pt: float = 16, bold: bool = False, italic: bool = False, color: Optional[RGBColor] = None):
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    # Set complex script font for Thai support in Word
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), FONT_NAME)
    rFonts.set(qn('w:hAnsi'), FONT_NAME)
    rFonts.set(qn('w:cs'), FONT_NAME)

def add_formatted_runs(p, text: str, default_size: float = 16, default_color: Optional[RGBColor] = None):
    # Parses basic markdown **bold** and *italic*
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = p.add_run(token[2:-2])
            set_run_font(run, size_pt=default_size, bold=True, color=default_color)
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            set_run_font(run, size_pt=default_size, italic=True, color=default_color)
        else:
            run = p.add_run(token)
            set_run_font(run, size_pt=default_size, bold=False, color=default_color)

def build_executive_brief_docx(title: str, date: str, lecturer: str, brief_markdown: str) -> io.BytesIO:
    doc = docx.Document()

    # Set 1-inch margins
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    # Document Title: Executive Brief (Centered, 18pt bold)
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(12)
    t_run = title_p.add_run("Executive Brief")
    set_run_font(t_run, size_pt=18, bold=True, color=RGBColor(15, 41, 66))

    # Metadata Table (3 rows, 2 cols)
    table = doc.add_table(rows=3, cols=2)
    table.autofit = False
    
    rows_data = [
        ("หัวข้อ", title or "-"),
        ("วันที่", date or "-"),
        ("อาจารย์/วิทยากร", lecturer or "-")
    ]

    for idx, (label, val) in enumerate(rows_data):
        row = table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(1.8)
        c1.width = Inches(4.7)

        # Left cell (Label)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(label)
        set_run_font(r0, size_pt=16, bold=True, color=RGBColor(15, 41, 66))

        # Right cell (Value)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(val)
        set_run_font(r1, size_pt=16, bold=False)

    # Style table borders
    tblPr = table._tbl.tblPr
    borders_xml = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_xml)

    # Space after table
    sep_p = doc.add_paragraph()
    sep_p.paragraph_format.space_before = Pt(0)
    sep_p.paragraph_format.space_after = Pt(8)

    # Process Brief Markdown
    lines = (brief_markdown or "").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue

        # Skip duplicate titles/headers in markdown body
        clean_check = line.replace('#', '').strip()
        if re.search(r'^(executive\s*brief|รายงานสรุป|หัวข้อ[:\s]|วันที่[:\s]|วิทยากร[:\s])', clean_check, re.IGNORECASE):
            continue

        # Check for Headings
        if line.startswith('#'):
            h_text = re.sub(r'^#+\s*', '', line).strip()
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = Pt(14)
            hp.paragraph_format.space_after = Pt(4)
            hp.paragraph_format.line_spacing = 1.15
            run = hp.add_run(h_text)
            set_run_font(run, size_pt=18, bold=True, color=RGBColor(15, 41, 66))
            continue

        # Check for Bullet Points
        if re.match(r'^[-*•]\s+', line) or re.match(r'^\d+\.\s+', line):
            # Bullet or numbered item
            b_match = re.match(r'^([-*•]|\d+\.)\s+(.*)', line)
            b_prefix = b_match.group(1) if b_match else '-'
            b_content = b_match.group(2) if b_match else line

            bp = doc.add_paragraph()
            bp.paragraph_format.left_indent = Inches(0.25)
            bp.paragraph_format.space_before = Pt(2)
            bp.paragraph_format.space_after = Pt(3)
            bp.paragraph_format.line_spacing = 1.15

            # Prefix run (e.g. • or 1.)
            bullet_char = "•  " if b_prefix in ['-', '*', '•'] else f"{b_prefix} "
            p_run = bp.add_run(bullet_char)
            set_run_font(p_run, size_pt=16, bold=True, color=RGBColor(23, 101, 203))

            add_formatted_runs(bp, b_content, default_size=16)
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_runs(p, line, default_size=16)

    # Save to in-memory buffer
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

def build_transcript_text(title: str, date: str, lecturer: str, transcript: str, notes: str = "", source_text: str = "") -> str:
    lines = []
    lines.append("=" * 64)
    lines.append("บันทึกคำบรรยายและถอดเสียง (Lecture Transcript)")
    lines.append("หลักสูตรวิทยาลัยป้องกันราชอาณาจักร (วปอ.) - ThaiNDC")
    lines.append("=" * 64)
    lines.append(f"หัวข้อ: {title or '-'}")
    lines.append(f"วันที่: {date or '-'}")
    lines.append(f"วิทยากร: {lecturer or '-'}")
    lines.append("=" * 64)
    lines.append("")
    if transcript and transcript.strip():
        lines.append("--- [1] บันทึกคำบรรยายถอดเสียง (Audio/Video Transcript) ---")
        lines.append(transcript.strip())
        lines.append("")
    if notes and notes.strip():
        lines.append("--- [2] บันทึกเพิ่มเติมของผู้เรียน (Personal Notes) ---")
        lines.append(notes.strip())
        lines.append("")
    if source_text and source_text.strip():
        lines.append("--- [3] เนื้อหาจากเอกสารประกอบการบรรยายและสไลด์ (Slides & Docs) ---")
        lines.append(source_text.strip())
        lines.append("")
    return "\n".join(lines)

def build_transcript_docx(title: str, date: str, lecturer: str, transcript: str, notes: str = "", source_text: str = "") -> io.BytesIO:
    doc = docx.Document()

    # Set 1-inch margins
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    # Document Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(12)
    t_run = title_p.add_run("บันทึกคำบรรยายและถอดเสียง (Lecture Transcript)")
    set_run_font(t_run, size_pt=18, bold=True, color=RGBColor(15, 41, 66))

    # Metadata Table
    table = doc.add_table(rows=3, cols=2)
    table.autofit = False
    rows_data = [
        ("หัวข้อ", title or "-"),
        ("วันที่", date or "-"),
        ("อาจารย์/วิทยากร", lecturer or "-")
    ]
    for idx, (label, val) in enumerate(rows_data):
        row = table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(1.8)
        c1.width = Inches(4.7)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(label)
        set_run_font(r0, size_pt=16, bold=True, color=RGBColor(15, 41, 66))

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(val)
        set_run_font(r1, size_pt=16, bold=False)

    tblPr = table._tbl.tblPr
    borders_xml = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_xml)

    sep_p = doc.add_paragraph()
    sep_p.paragraph_format.space_before = Pt(0)
    sep_p.paragraph_format.space_after = Pt(8)

    def add_section(header_title: str, content: str):
        if not content or not content.strip():
            return
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(14)
        hp.paragraph_format.space_after = Pt(4)
        hp.paragraph_format.line_spacing = 1.15
        hrun = hp.add_run(header_title)
        set_run_font(hrun, size_pt=17, bold=True, color=RGBColor(15, 41, 66))

        for block in content.strip().split('\n\n'):
            block = block.strip()
            if not block:
                continue
            bp = doc.add_paragraph()
            bp.paragraph_format.space_before = Pt(2)
            bp.paragraph_format.space_after = Pt(4)
            bp.paragraph_format.line_spacing = 1.15
            run = bp.add_run(block)
            set_run_font(run, size_pt=16, bold=False)

    if transcript:
        add_section("🎙️ บันทึกคำบรรยายถอดเสียง (Audio/Video Transcript)", transcript)
    if notes:
        add_section("📝 บันทึกเพิ่มเติมของผู้เรียน (Personal Notes)", notes)
    if source_text:
        add_section("📑 เนื้อหาจากเอกสารและสไลด์ประกอบการบรรยาย (Slides & Documents)", source_text)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

