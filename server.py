"""WPOI Learning Brief: local prototype, not production deployment."""
import os, sqlite3, json, uuid, tempfile, time, re
from pathlib import Path
from datetime import datetime, timezone
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pptx import Presentation
from pypdf import PdfReader
from google import genai
from google.genai import types

ROOT=Path(__file__).resolve().parent
DATA=ROOT/'data'; DATA.mkdir(exist_ok=True)
DB=DATA/'lectures.sqlite3'
app=FastAPI(title='ThaiNDC AI Learning Assistant')
app.mount('/static',StaticFiles(directory=ROOT/'static'),name='static')

def conn():
 c=sqlite3.connect(DB); c.row_factory=sqlite3.Row
 c.execute('''CREATE TABLE IF NOT EXISTS lectures (id TEXT PRIMARY KEY, date TEXT NOT NULL, title TEXT NOT NULL, lecturer TEXT, notes TEXT, source_text TEXT, transcript TEXT, brief TEXT, created_at TEXT NOT NULL)''')
 return c

def ai():
 key=os.getenv('GEMINI_API_KEY') or os.getenv('GOOGLE_API_KEY')
 if not key: raise HTTPException(400,'กรุณาตั้งค่า GEMINI_API_KEY ที่ฝั่งเซิร์ฟเวอร์')
 return genai.Client(api_key=key)

AUDIO_MIME={'.mp3':'audio/mpeg','.mpeg':'audio/mpeg','.mpga':'audio/mpeg','.mp4':'audio/mp4','.m4a':'audio/mp4','.wav':'audio/wav','.webm':'audio/webm'}

def transcribe(path:str, suffix:str)->str:
 c=ai(); up=c.files.upload(file=path,config={'mime_type':AUDIO_MIME[suffix]})
 try:
  while up.state and up.state.name=='PROCESSING': time.sleep(1); up=c.files.get(name=up.name)
  models=[os.getenv('TRANSCRIBE_MODEL','gemini-2.5-flash'),'gemini-2.0-flash','gemini-1.5-flash']
  prompt='ถอดเสียงไฟล์นี้เป็นข้อความตามที่พูดจริงในภาษาเดิม ห้ามสรุป ห้ามแต่งเติม ถ้าฟังไม่ชัดให้ใส่ [ไม่ชัด] ตอบเฉพาะข้อความถอดเสียง'
  last_err=None
  for m in models:
   for attempt in range(2):
    try:
     out=c.models.generate_content(model=m,contents=[up,prompt])
     return out.text or ''
    except Exception as e:
     last_err=e
     time.sleep(2)
  raise HTTPException(503,f'AI โมเดลไม่พร้อมใช้งานชั่วคราว ({str(last_err)}) กรุณาลองใหม่อีกครั้ง')
 finally:
  try: c.files.delete(name=up.name)
  except Exception: pass


def extract(name: str, raw: bytes):
 suffix=Path(name).suffix.lower()
 with tempfile.NamedTemporaryFile(suffix=suffix,delete=False) as f:
  f.write(raw); p=f.name
 try:
  if suffix=='.pptx':
   pres=Presentation(p)
   return '\n\n'.join(f'[สไลด์ {i}] '+ ' | '.join(sh.text for sh in slide.shapes if sh.has_text_frame and sh.text.strip()) for i,slide in enumerate(pres.slides,1))
  if suffix=='.pdf':
   reader=PdfReader(p)
   return '\n\n'.join(f'[หน้า {i}] '+(page.extract_text() or '') for i,page in enumerate(reader.pages,1))
  if suffix=='.txt': return raw.decode('utf-8-sig',errors='replace')
  raise HTTPException(400,'รองรับไฟล์สไลด์ .pptx, .pdf, .txt')
 finally: os.remove(p)

@app.get('/')
def index(): return FileResponse(ROOT/'static'/'index.html')

@app.get('/api/lectures')
def list_lectures():
 with conn() as c:
  return [dict(r) for r in c.execute('SELECT id,date,title,lecturer,created_at,CASE WHEN brief IS NULL THEN 0 ELSE 1 END AS has_brief FROM lectures ORDER BY date DESC,created_at DESC')]

def format_brief_html(title: str, date: str, lecturer: str, md_text: str) -> str:
 if not md_text: return ''
 style_block='''<style>
div, p, td, li { font-family: 'TH SarabunPSK', 'TH Sarabun New', sans-serif; font-size: 16pt; }
h1 { font-family: 'TH SarabunPSK', 'TH Sarabun New', sans-serif; font-size: 18pt; font-weight: bold; margin-top: 12pt; margin-bottom: 4pt; mso-margin-top-alt: 12pt; mso-margin-bottom-alt: 4pt; color: #000; }
table { border-collapse: collapse; width: 100%; margin-bottom: 14pt; border: 1px solid #000; }
table td { padding: 3pt 6pt; border: 1px solid #000; vertical-align: middle; }
table td p { margin: 0pt !important; margin-top: 0pt !important; margin-bottom: 0pt !important; mso-margin-top-alt: 0pt !important; mso-margin-bottom-alt: 0pt !important; line-height: 1.15; font-family: 'TH SarabunPSK', 'TH Sarabun New', sans-serif; font-size: 16pt; }
</style>'''
 html=[style_block, '''<div style="font-family: 'TH SarabunPSK', 'TH Sarabun New', sans-serif; font-size: 16pt; line-height: 1.25; color: #000;">''']
 html.append('''<p style="text-align: center; font-size: 18pt; font-weight: bold; margin-top: 0; margin-bottom: 12pt; mso-margin-top-alt: 0pt; mso-margin-bottom-alt: 12pt;">Executive Brief</p>''')
 html.append(f'''<table border="1" style="border-collapse: collapse; width: 100%; font-size: 16pt; margin-bottom: 14pt; border: 1px solid #000;">
  <tr>
    <td style="width: 25%; padding: 3pt 6pt; border: 1px solid #000; vertical-align: middle;"><p style="margin: 0pt; margin-top: 0pt; margin-bottom: 0pt; mso-margin-top-alt: 0pt; mso-margin-bottom-alt: 0pt; line-height: 1.15; font-weight: bold; font-family: 'TH SarabunPSK', 'TH Sarabun New', sans-serif; font-size: 16pt;">หัวข้อ</p></td>
    <td style="padding: 3pt 6pt; border: 1px solid #000; vertical-align: middle;"><p style="margin: 0pt; margin-top: 0pt; margin-bottom: 0pt; mso-margin-top-alt: 0pt; mso-margin-bottom-alt: 0pt; line-height: 1.15; font-family: 'TH SarabunPSK', 'TH Sarabun New', sans-serif; font-size: 16pt;">{title or '-'}</p></td>
  </tr>
  <tr>
    <td style="padding: 3pt 6pt; border: 1px solid #000; vertical-align: middle;"><p style="margin: 0pt; margin-top: 0pt; margin-bottom: 0pt; mso-margin-top-alt: 0pt; mso-margin-bottom-alt: 0pt; line-height: 1.15; font-weight: bold; font-family: 'TH SarabunPSK', 'TH Sarabun New', sans-serif; font-size: 16pt;">วันที่</p></td>
    <td style="padding: 3pt 6pt; border: 1px solid #000; vertical-align: middle;"><p style="margin: 0pt; margin-top: 0pt; margin-bottom: 0pt; mso-margin-top-alt: 0pt; mso-margin-bottom-alt: 0pt; line-height: 1.15; font-family: 'TH SarabunPSK', 'TH Sarabun New', sans-serif; font-size: 16pt;">{date or '-'}</p></td>
  </tr>
  <tr>
    <td style="padding: 3pt 6pt; border: 1px solid #000; vertical-align: middle;"><p style="margin: 0pt; margin-top: 0pt; margin-bottom: 0pt; mso-margin-top-alt: 0pt; mso-margin-bottom-alt: 0pt; line-height: 1.15; font-weight: bold; font-family: 'TH SarabunPSK', 'TH Sarabun New', sans-serif; font-size: 16pt;">อาจารย์/วิทยากร</p></td>
    <td style="padding: 3pt 6pt; border: 1px solid #000; vertical-align: middle;"><p style="margin: 0pt; margin-top: 0pt; margin-bottom: 0pt; mso-margin-top-alt: 0pt; mso-margin-bottom-alt: 0pt; line-height: 1.15; font-family: 'TH SarabunPSK', 'TH Sarabun New', sans-serif; font-size: 16pt;">{lecturer or '-'}</p></td>
  </tr>
</table>''')
 in_list=False
 past_header=False
 section_keywords=['executive summary','key takeaways','strategic insights','implications for leaders','evidence & source references','reflection questions','ข้อจำกัดของข้อมูล']
 for raw in md_text.strip().split('\n'):
  line=raw.strip()
  if not line:
   if in_list: html.append('</ul>'); in_list=False
   continue
  clean_lower=line.lower()
  if not past_header:
   if any(kw in clean_lower for kw in ['รายงานสรุป', 'executive brief', 'หัวข้อ:', 'วันที่:', 'วิทยากร:', 'หัวข้อ :', 'วันที่ :', 'วิทยากร :']) \
      or (line.startswith('*') and any(kw in clean_lower for kw in ['หัวข้อ', 'วันที่', 'วิทยากร'])) \
      or (line.startswith('-') and any(kw in clean_lower for kw in ['หัวข้อ', 'วันที่', 'วิทยากร'])):
    continue
   if 'executive summary' in clean_lower or line.startswith('#'):
    past_header=True
  line_clean=re.sub(r'^\[.*?\]\s*','',line)
  line_formatted=re.sub(r'\*\*(.+?)\*\*',r'<b>\1</b>',line_clean)
  if re.match(r'^#{1,6}\s+',line):
   if in_list: html.append('</ul>'); in_list=False
   h_text=re.sub(r'^#{1,6}\s+','',line_formatted).strip()
   html.append(f'''<h1 style="font-family: 'TH SarabunPSK', 'TH Sarabun New', sans-serif; font-size: 18pt; font-weight: bold; margin-top: 12pt; margin-bottom: 4pt; mso-margin-top-alt: 12pt; mso-margin-bottom-alt: 4pt; color: #000;">{h_text}</h1>''')
   continue
  bullet_match=re.match(r'^(\*|\-|\•|\d+\.)\s*(.*)$',line_clean)
  if bullet_match:
   if not in_list: html.append('''<ul style="margin-top: 0; margin-bottom: 6pt; padding-left: 20pt;">'''); in_list=True
   item_raw=bullet_match.group(2)
   item_formatted=re.sub(r'\*\*(.+?)\*\*',r'<b>\1</b>',item_raw)
   html.append(f'''<li style="font-family: 'TH SarabunPSK', 'TH Sarabun New', sans-serif; font-size: 16pt; margin-bottom: 4pt; text-align: justify; text-justify: inter-cluster;">{item_formatted}</li>''')
   continue
  clean_title_check=re.sub(r'^[#\*\s]+','',line_clean).lower()
  if any(kw in clean_title_check for kw in section_keywords) and len(line_clean)<70:
   if in_list: html.append('</ul>'); in_list=False
   clean_title=re.sub(r'^[\*\s]+|[\*\s]+$','',line_formatted)
   html.append(f'''<h1 style="font-family: 'TH SarabunPSK', 'TH Sarabun New', sans-serif; font-size: 18pt; font-weight: bold; margin-top: 12pt; margin-bottom: 4pt; mso-margin-top-alt: 12pt; mso-margin-bottom-alt: 4pt; color: #000;">{clean_title}</h1>''')
   continue
  if in_list: html.append('</ul>'); in_list=False
  html.append(f'''<p style="font-family: 'TH SarabunPSK', 'TH Sarabun New', sans-serif; font-size: 16pt; margin-top: 0; margin-bottom: 6pt; text-align: justify; text-justify: inter-cluster;">{line_formatted}</p>''')
 if in_list: html.append('</ul>')
 html.append('</div>')
 return '\n'.join(html)

@app.get('/api/lectures/{id}')
def get_lecture(id:str):
 with conn() as c: row=c.execute('SELECT * FROM lectures WHERE id=?',(id,)).fetchone()
 if not row: raise HTTPException(404,'ไม่พบรายการ')
 d=dict(row)
 d['brief_html']=format_brief_html(d.get('title') or '',d.get('date') or '',d.get('lecturer') or '',d.get('brief') or '')
 return d

@app.post('/api/lectures')
async def create_lecture(date:str=Form(...),title:str=Form(...),lecturer:str=Form(''),notes:str=Form(''), slides:list[UploadFile]|None=File(None),audio:list[UploadFile]|None=File(None)):
 if not title.strip() or not date.strip(): raise HTTPException(400,'โปรดระบุวันที่และหัวข้อ')
 parts=[]
 for slide in [f for f in slides or [] if f.filename]:
  raw=await slide.read()
  if len(raw)>20*1024*1024: raise HTTPException(413,'ไฟล์สไลด์แต่ละไฟล์ต้องไม่เกิน 20MB')
  parts.append(f'=== เอกสาร {slide.filename} ===\n'+extract(slide.filename or '',raw))
 transcript_parts=[]
 for idx,sound in enumerate([f for f in audio or [] if f.filename],1):
  raw=await sound.read()
  if len(raw)>25*1024*1024: raise HTTPException(413,'ไฟล์เสียงแต่ละไฟล์ต้องไม่เกิน 25MB; กรุณาแบ่งไฟล์ก่อนอัปโหลด')
  suffix=Path(sound.filename or '').suffix.lower()
  if suffix not in {'.mp3','.mp4','.mpeg','.mpga','.m4a','.wav','.webm'}: raise HTTPException(400,'ไฟล์เสียงไม่รองรับ')
  with tempfile.NamedTemporaryFile(suffix=suffix,delete=False) as f: f.write(raw); p=f.name
  try:
   transcript_parts.append(f'[เสียงชุด {idx}: {sound.filename}]\n{transcribe(p,suffix)}')
  finally: os.remove(p)
 id=str(uuid.uuid4()); now=datetime.now(timezone.utc).isoformat()
 with conn() as c:
  c.execute('INSERT INTO lectures VALUES(?,?,?,?,?,?,?,?,?)',(id,date,title,lecturer,notes,'\n'.join(parts),'\n\n'.join(transcript_parts),None,now))
 return {'id':id,'status':'saved','slide_count':len(parts),'audio_count':len(transcript_parts)}

@app.post('/api/lectures/{id}/summarize')
def summarize(id:str):
 row=get_lecture(id)
 source=f"วันที่: {row['date']}\nหัวข้อ: {row['title']}\nวิทยากร: {row['lecturer']}\nบันทึกผู้เรียน: {row['notes']}\n\nสไลด์:\n{row['source_text']}\n\nถอดเสียง:\n{row['transcript']}"
 if not any((row['notes'],row['source_text'],row['transcript'])): raise HTTPException(400,'ยังไม่มีเนื้อหาสำหรับสรุป')
 # MVP safeguard: refuse silent truncation, so long lectures require a chunking worker in production.
 if len(source)>100000: raise HTTPException(413,'เนื้อหายาวเกินขีดจำกัดต้นแบบ กรุณาแบ่งการบรรยายเป็นหลายรายการ')
 models=[os.getenv('BRIEF_MODEL','gemini-2.5-flash'),'gemini-2.0-flash','gemini-1.5-flash']
 client=ai(); out=None; last_err=None
 sys_instruction='''คุณเป็นเลขานุการวิชาการสำหรับผู้บริหารระดับสูงของหลักสูตร วปอ. (ThaiNDC). เขียนภาษาไทยกึ่งทางการ กระชับ ตรงประเด็น ห้ามแต่งข้อเท็จจริง/คำพูด/เวลาจากเสียงที่ไม่มี timecode ให้แยก "ข้อเท็จจริงจากแหล่งข้อมูล" และ "การวิเคราะห์ต่อยอดของ AI" อย่างชัดเจน ใช้แหล่งอ้างอิง [สไลด์ N]/[หน้า N]/[เสียงชุด N] เฉพาะเมื่อมีหลักฐาน อย่าอ้างแหล่งเท็จ หากมีความขัดแย้งให้ระบุว่าไม่สอดคล้องกัน ไม่ถ่ายทอดข้อมูลอ่อนไหวที่ไม่จำเป็น จัดรูปแบบ Markdown มี: Executive Summary, Key Takeaways 5-7 ข้อ, Strategic Insights, Implications for Leaders, Evidence & Source References, Reflection Questions 3-5 ข้อ, ข้อจำกัดของข้อมูล. ข้อสำคัญ: ห้ามพิมพ์ชื่อเรื่อง วันที่ หรือชื่อวิทยากรซ้ำที่ส่วนต้น (เนื่องจากระบบมีตารางหัวกระดาษให้อยู่แล้ว) ให้เริ่มต้นเนื้อหาด้วยหัวข้อ Executive Summary ทันที.'''
 for m in models:
  for attempt in range(2):
   try:
    out=client.models.generate_content(model=m,config=types.GenerateContentConfig(system_instruction=sys_instruction),contents=source)
    if out and out.text: break
   except Exception as e:
    last_err=e
    time.sleep(2)
  if out and out.text: break
 if not out or not out.text: raise HTTPException(503,f'AI โมเดลไม่พร้อมใช้งานชั่วคราว ({str(last_err)})')
 brief=out.text
 with conn() as c: c.execute('UPDATE lectures SET brief=? WHERE id=?',(brief,id))
 brief_html=format_brief_html(row.get('title') or '',row.get('date') or '',row.get('lecturer') or '',brief)
 return {'id':id,'brief':brief,'brief_html':brief_html}
